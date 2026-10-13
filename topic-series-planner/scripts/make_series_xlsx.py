# -*- coding: utf-8 -*-
"""把「主题 -> 合集 -> 各集标题+时长」生成为 xls 清单。

用法：
  python -X utf8 make_series_xlsx.py <series.json> [--out 输出.xlsx]

输入 JSON 结构（见 references/series-schema.md）：
{
  "topic": "宇宙科普",
  "series": [
    {
      "name": "火星探索",
      "position": "人类下一个家园，也是最大的谜团",
      "episodes": [
        {"title": "火星上的一天比地球多40分钟，藏着什么秘密", "angle": "时间", "hook": "反常识", "minutes": 5},
        ...
      ]
    },
    ...
  ]
}

校验（任一失败即报错，不生成文件）：
  1. 全局：所有集的标题文本两两不同（去重）
  2. 合集内：所有集的"角度"两两不同（不撞选题）
  3. 每集 minutes 必须是 1-60 的整数

输出 xlsx 含两个表：
  - <主题>清单：序号/合集/集数/标题/选题角度/钩子类型/时长(分钟)/累计时长
  - 合集概览：合集/集数/总时长/一句话定位
"""
import json, sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def validate(data):
    topic = data.get("topic", "主题")
    series = data.get("series", [])
    if not series:
        raise SystemExit("series 为空")
    all_titles, problems = [], []
    total_ep = 0
    for s in series:
        name = s.get("name", "?")
        eps = s.get("episodes", [])
        if not eps:
            problems.append("合集【" + name + "】没有集")
        angles = []
        for e in eps:
            t = (e.get("title") or "").strip()
            a = (e.get("angle") or "").strip()
            m = e.get("minutes")
            if not t:
                problems.append("合集【" + name + "】有集缺 title")
            all_titles.append(t)
            angles.append(a)
            if not isinstance(m, int) or m < 1 or m > 60:
                problems.append("《" + t + "》时长非法（应为 1-60 整数分钟），实际=" + str(m))
            total_ep += 1
        if len(angles) != len(set(angles)):
            dup = [x for x in set(angles) if angles.count(x) > 1]
            problems.append("合集【" + name + "】选题角度重复：" + "/".join(dup))
    dups = [t for t in set(all_titles) if all_titles.count(t) > 1]
    if dups:
        problems.append("标题文本重复：" + "；".join(dups[:5]) + ("…" if len(dups) > 5 else ""))
    if problems:
        print("校验未通过，不生成文件：")
        for p in problems:
            print("  ✗ " + p)
        sys.exit(1)
    return topic, series, total_ep


def build(topic, series, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (topic[:24] or "清单")
    hdr_fill = PatternFill("solid", fgColor="1F2937")
    hdr_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    theme_fills = ["2D3748", "4A5568"]
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["序号", "合集主题", "集数", "视频标题", "选题角度", "钩子类型", "时长(分钟)", "累计时长"]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c); cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = center; cell.border = border

    idx = 0
    for ti, s in enumerate(series):
        fill = PatternFill("solid", fgColor=theme_fills[ti % 2])
        cum = 0
        for n, e in enumerate(s["episodes"], 1):
            idx += 1
            cum += e["minutes"]
            ws.append([idx, s["name"], "第%d集" % n, e["title"], e.get("angle", ""), e.get("hook", ""), e["minutes"], cum])
            r = ws.max_row
            for c in range(1, 9):
                cell = ws.cell(r, c); cell.border = border
                cell.alignment = center if c in (1, 3, 6, 7, 8) else left
                if c == 2:
                    cell.fill = fill; cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True, size=10)
                elif c == 4:
                    cell.font = Font(name="微软雅黑", size=10)
                elif c == 5:
                    cell.font = Font(name="微软雅黑", size=10, color="2F855A", bold=True)
                elif c in (7, 8):
                    cell.font = Font(name="微软雅黑", size=10, color="2B6CB0", bold=True)
                else:
                    cell.font = Font(name="微软雅黑", size=10, color="555555")

    for col, w in [("A",6),("B",14),("C",8),("D",48),("E",12),("F",12),("G",11),("H",11)]:
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # 概览
    ws2 = wb.create_sheet("合集概览")
    for c, h in enumerate(["合集主题", "集数", "总时长(分钟)", "一句话定位"], 1):
        cell = ws2.cell(1, c); cell.value = h
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = center; cell.border = border
    for s in series:
        tot = sum(e["minutes"] for e in s["episodes"])
        ws2.append([s["name"], len(s["episodes"]), tot, s.get("position", "")])
        r = ws2.max_row
        for c in range(1, 5):
            ws2.cell(r, c).border = border
            ws2.cell(r, c).font = Font(name="微软雅黑", size=10)
            ws2.cell(r, c).alignment = center if c in (2, 3) else left
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 8
    ws2.column_dimensions["C"].width = 13
    ws2.column_dimensions["D"].width = 44

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    return out


def main():
    if len(sys.argv) < 2:
        print("用法: python -X utf8 make_series_xlsx.py <series.json> [--out 输出.xlsx]")
        sys.exit(1)
    src = sys.argv[1]
    data = load(src)
    topic, series, total = validate(data)
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else os.path.join(os.getcwd(), topic + "视频清单.xlsx")
    build(topic, series, out)
    tot_min = sum(e["minutes"] for s in series for e in s["episodes"])
    print("OK " + out)
    print("  主题: %s | 合集: %d | 总集数: %d | 总时长: %d 分钟" % (topic, len(series), total, tot_min))
    print("  去重校验：标题全唯一 + 各合集选题角度不重复 ✓")


if __name__ == "__main__":
    main()
